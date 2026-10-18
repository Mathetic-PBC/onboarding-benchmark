"""Shared helpers for the Phase 2 synthetic generators (_scripts/synth/gen_<tool>.py). Everything is generated from
synthetic/PERSON.md.

Run generators with the scratch venv python (python-docx, openpyxl, python-pptx, reportlab), e.g.
  $SYNTHPY _scripts/synth/gen_notion.py
where SYNTHPY=/private/tmp/claude-501/-Users-divadbaroon-Desktop-benchmark/ddaa4403-0276-46d5-8b1f-7fd3346e76de/scratchpad/synthvenv/bin/python

Each generator:
  g = Gen('notion')                       # synthetic/notion/, rebuilt from scratch on every run (deterministic)
  g.text('Workspace/Page 1a2b….md', body, created='2026-01-28 23:10', modified='2026-02-03 10:02',
         account='David's Notion', project_ids=['project_engelbart'], modeled_on='', notes='HARD CASE: vague title')
  ...
  g.finish()                              # writes MANIFEST.csv and _scripts/synth/labels/notion.json

Dates: strings 'YYYY-MM-DD HH:MM[:SS]' are local time where David was then (America/New_York until 2026-08-08, then
America/Los_Angeles), unless they carry an offset/Z. File times: birthtime = created, mtime = modified (on APFS, setting an
mtime older than the birthtime moves the birthtime back, so we set created first, then modified).

Labels: one per *item* (what build_source_items.py yields as an item key: a file's path relative to the tool folder, or a
directory for multi-file items such as a Gmail thread `<acct>/threads/<id>`, a Zoom meeting folder, a Claude project, a repo).
Pass label_key= on any file of a multi-file item (or call g.label()). Notes in labels and the manifest are hidden from the agent.
"""
import csv, datetime, email.utils, hashlib, io, json, os, random, shutil, sqlite3, subprocess, uuid, zipfile
from zoneinfo import ZoneInfo
from email.message import EmailMessage

ROOT = os.path.expanduser('~/Desktop/benchmark/benchmark-data')
SYN = os.path.join(ROOT, 'synthetic')
LABEL_DIR = os.path.join(ROOT, '_scripts', 'synth', 'labels')
FIELDS = ['path', 'tool', 'account_or_library', 'date', 'origin', 'source', 'modeled_on', 'notes']
REF = datetime.datetime(2026, 10, 17, 23, 59, tzinfo=ZoneInfo('America/Los_Angeles'))
START = datetime.datetime(2025, 10, 17, 0, 0, tzinfo=ZoneInfo('America/New_York'))
ET, PT, UTC = ZoneInfo('America/New_York'), ZoneInfo('America/Los_Angeles'), datetime.timezone.utc
MOVE = datetime.datetime(2026, 8, 8, 12, 0, tzinfo=ET)
ME = ('David Barron', 'dbarron410@vt.edu')
ME_GMAIL = ('David Barron', 'davidbarron410@gmail.com')
ME_RUNS = ('David Barron', 'david.barron.runs@gmail.com')
PROJECTS = {r['id'] for r in json.load(open(os.path.join(ROOT, 'ground_truth_projects.json')))} | {'project_hint_timing_study'}

# ---------------- dates ----------------

def tz_for(naive_or_aware):
    d = naive_or_aware
    probe = d if d.tzinfo else d.replace(tzinfo=ET)
    return PT if probe >= MOVE else ET


def dt(s):
    """'2026-03-31 23:12' (local to where David was) | ISO with offset | datetime -> aware datetime."""
    if isinstance(s, datetime.datetime):
        return s if s.tzinfo else s.replace(tzinfo=tz_for(s))
    s = s.strip().replace('T', ' ')
    if s.endswith('Z'): s = s[:-1] + '+00:00'
    try:
        d = datetime.datetime.fromisoformat(s)
    except ValueError:
        d = datetime.datetime.strptime(s, '%Y-%m-%d')
    return d if d.tzinfo else d.replace(tzinfo=tz_for(d))


def iso(d, z=False):
    d = dt(d)
    return d.astimezone(UTC).isoformat(timespec='seconds').replace('+00:00', 'Z') if z else d.isoformat(timespec='seconds')


def epoch(d):
    return dt(d).timestamp()


def rfc2822(d):
    return email.utils.format_datetime(dt(d))


# PERSON.md §4 burst calendar: (week start, weeks covered, weight)
BURSTS = [('2025-10-20', 1, 1.0), ('2025-11-03', 1, 1.8), ('2025-11-10', 1, 2.5), ('2025-12-08', 1, 2.2), ('2025-12-22', 2, 0.2),
          ('2026-01-12', 1, 1.4), ('2026-01-26', 1, 2.0), ('2026-02-09', 1, 1.6), ('2026-02-16', 1, 0.3), ('2026-03-23', 1, 2.6),
          ('2026-03-30', 1, 3.0), ('2026-04-06', 1, 0.4), ('2026-04-27', 1, 1.8), ('2026-05-04', 1, 2.8), ('2026-05-18', 2, 0.3),
          ('2026-06-15', 1, 2.4), ('2026-06-29', 1, 0.2), ('2026-07-27', 1, 2.4), ('2026-08-03', 1, 2.0), ('2026-08-10', 1, 2.6),
          ('2026-08-24', 1, 2.2), ('2026-09-07', 1, 2.6), ('2026-09-21', 1, 1.8), ('2026-10-05', 1, 3.0), ('2026-10-12', 1, 2.4)]


def week_weights(start=None, end=None):
    """[(monday date, weight)] for every week from start to end (defaults: the scope year)."""
    s = dt(start or '2025-10-13').date(); e = dt(end or '2026-10-17').date()
    s -= datetime.timedelta(days=s.weekday())
    table = {}
    for w, n, wt in BURSTS:
        d0 = datetime.date.fromisoformat(w)
        for i in range(n): table[d0 + datetime.timedelta(weeks=i)] = wt
    out = []
    while s <= e:
        out.append((s, table.get(s, 0.8))); s += datetime.timedelta(weeks=1)
    return out


DAY_W = [1.3, 1.2, 1.15, 1.0, 0.8, 0.6, 1.1]  # Mon..Sun
HOUR_W = [0.5, 0.35, 0.2, 0.1, 0.05, 0.05, 0.1, 0.4, 0.7, 0.9, 1.0, 1.0, 0.8, 0.9, 1.0, 1.0, 0.9, 0.8, 0.7, 0.8, 0.9, 1.0, 1.0, 0.8]


def bursty_dates(n, seed, start=None, end=None, extra_weight=None):
    """n sorted aware datetimes inside [start, end] following the PERSON.md burst calendar. extra_weight(monday)->float
    multiplies a week's weight (e.g. to concentrate a project's items in its own weeks)."""
    rnd = random.Random(seed)
    lo = dt(start or START); hi = min(dt(end or REF), REF)
    weeks = [(m, w * (extra_weight(m) if extra_weight else 1.0)) for m, w in week_weights(lo, hi)]
    out = []
    while len(out) < n:
        m = rnd.choices([w[0] for w in weeks], weights=[w[1] for w in weeks])[0]
        day = m + datetime.timedelta(days=rnd.choices(range(7), weights=DAY_W)[0])
        hour = rnd.choices(range(24), weights=HOUR_W)[0]
        naive = datetime.datetime(day.year, day.month, day.day, hour, rnd.randrange(60), rnd.randrange(60))
        d = naive.replace(tzinfo=tz_for(naive))
        if lo <= d <= hi: out.append(d)
    return sorted(out)


def later(d, minutes=0, hours=0, days=0, seed=None, jitter_min=0):
    r = random.Random(seed).randint(0, jitter_min) if jitter_min else 0
    return dt(d) + datetime.timedelta(minutes=minutes + r, hours=hours, days=days)


def uid(seed):
    """Deterministic UUID4-looking id from a seed string."""
    h = hashlib.sha256(str(seed).encode()).hexdigest()
    return f'{h[:8]}-{h[8:12]}-4{h[13:16]}-{"89ab"[int(h[16], 16) % 4]}{h[17:20]}-{h[20:32]}'


def hexid(seed, n=32):
    return hashlib.sha256(str(seed).encode()).hexdigest()[:n]

# ---------------- generator ----------------

class Gen:
    def __init__(self, tool, clean=True):
        self.tool = tool
        self.dir = os.path.join(SYN, tool)
        if clean and os.path.isdir(self.dir):
            for x in os.listdir(self.dir):
                if x == 'source_items.json': continue
                p = os.path.join(self.dir, x)
                shutil.rmtree(p) if os.path.isdir(p) and not os.path.islink(p) else os.remove(p)
        os.makedirs(self.dir, exist_ok=True)
        self.rows = {}
        self.labels = {}

    # -- bookkeeping
    def p(self, rel):
        out = os.path.join(self.dir, rel)
        os.makedirs(os.path.dirname(out), exist_ok=True)
        return out

    def times(self, path, created, modified=None):
        c = epoch(created); m = epoch(modified or created)
        os.utime(path, (c, c))
        os.utime(path, (m, m))

    def add(self, rel, *, account, date, project_ids=None, modeled_on='', notes='', source='', label_key=None, label_notes=None):
        """Manifest row for one file (+ label for its item)."""
        if project_ids is not None:
            bad = set(project_ids) - PROJECTS
            assert not bad, f'unknown project ids {bad}'
        self.rows[rel] = dict(path=rel, tool=self.tool, account_or_library=account, date=iso(date) if date else '', origin='synthetic',
                              source=source, modeled_on=modeled_on, notes=notes)
        if project_ids is not None:
            self.label(label_key or rel, project_ids, label_notes if label_notes is not None else (notes or None))

    def label(self, key, project_ids, notes=None):
        bad = set(project_ids) - PROJECTS
        assert not bad, f'unknown project ids {bad}'
        old = self.labels.get(key)
        if old:  # several files of one item: union of projects, first notes win
            old['project_ids'] = sorted(set(old['project_ids']) | set(project_ids))
            if notes and not old.get('notes'): old['notes'] = notes
        else:
            self.labels[key] = {'project_ids': sorted(set(project_ids)), 'notes': notes}

    def _done(self, rel, created, modified, meta):
        path = os.path.join(self.dir, rel)
        self.times(path, created, modified)
        self.add(rel, date=modified or created, **meta)
        return path

    # -- writers (meta = account, project_ids, modeled_on, notes, source, label_key, label_notes)
    def text(self, rel, content, created, modified=None, **meta):
        with open(self.p(rel), 'w', encoding='utf-8') as f: f.write(content)
        return self._done(rel, created, modified, meta)

    def bytes(self, rel, data, created, modified=None, **meta):
        with open(self.p(rel), 'wb') as f: f.write(data)
        return self._done(rel, created, modified, meta)

    def json(self, rel, obj, created, modified=None, indent=1, **meta):
        return self.text(rel, json.dumps(obj, indent=indent, ensure_ascii=False), created, modified, **meta)

    def jsonl(self, rel, objs, created, modified=None, **meta):
        return self.text(rel, ''.join(json.dumps(o, ensure_ascii=False) + '\n' for o in objs), created, modified, **meta)

    def docx(self, rel, blocks, created, modified=None, title=None, author='David Barron', **meta):
        return self.bytes(rel, docx_bytes(blocks, title, author, created, modified), created, modified, **meta)

    def xlsx(self, rel, sheets, created, modified=None, title=None, author='David Barron', **meta):
        return self.bytes(rel, xlsx_bytes(sheets, title, author, created, modified), created, modified, **meta)

    def pptx(self, rel, slides, created, modified=None, title=None, author='David Barron', **meta):
        return self.bytes(rel, pptx_bytes(slides, title, author, created, modified), created, modified, **meta)

    def pdf(self, rel, blocks, created, modified=None, title=None, author='David Barron', **meta):
        return self.bytes(rel, pdf_bytes(blocks, title, author, created, modified), created, modified, **meta)

    def eml(self, rel, created=None, **kw):
        """kw: frm, to, subject, date, body, cc, html, attachments=[(name, bytes, 'type/subtype')], message_id, in_reply_to,
        references, headers={...}; plus manifest meta. created defaults to the Date."""
        meta = {k: kw.pop(k) for k in list(kw) if k in ('account', 'project_ids', 'modeled_on', 'notes', 'source', 'label_key', 'label_notes')}
        data = eml_bytes(**kw)
        return self.bytes(rel, data, created or kw['date'], kw['date'], **meta)

    def zip(self, rel, files, created, modified=None, **meta):
        """files: {name: bytes|str} or [(name, bytes|str, date)]"""
        buf = io.BytesIO()
        items = files.items() if isinstance(files, dict) else [(n, b) for n, b, *_ in files]
        dates = {} if isinstance(files, dict) else {n: d for n, b, *rest in files for d in rest[:1]}
        with zipfile.ZipFile(buf, 'w', zipfile.ZIP_DEFLATED) as z:
            for n, b in items:
                d = dt(dates.get(n) or modified or created)
                zi = zipfile.ZipInfo(n, date_time=d.timetuple()[:6]); zi.compress_type = zipfile.ZIP_DEFLATED
                z.writestr(zi, b if isinstance(b, (bytes, bytearray)) else b.encode())
        return self.bytes(rel, buf.getvalue(), created, modified, **meta)

    def copy_from(self, rel, src, created, modified=None, **meta):
        """Copy a file from another synthetic folder / _scripts/synth/shared (never from authentic/)."""
        assert '/authentic/' not in os.path.abspath(src)
        shutil.copyfile(src, self.p(rel))
        return self._done(rel, created, modified, meta)

    def sqlite(self, rel, schema_sql, rows_by_table, created, modified=None, **meta):
        """rows_by_table: {table: [tuple,...]} inserted with ? placeholders."""
        path = self.p(rel)
        if os.path.exists(path): os.remove(path)
        con = sqlite3.connect(path); con.executescript(schema_sql)
        for t, rows in rows_by_table.items():
            if rows: con.executemany(f'INSERT INTO {t} VALUES ({",".join("?" * len(rows[0]))})', rows)
        con.commit(); con.close()
        return self._done(rel, created, modified, meta)

    def dir_times(self, rel, created, modified=None):
        self.times(os.path.join(self.dir, rel), created, modified)

    def finish(self):
        with open(os.path.join(self.dir, 'MANIFEST.csv'), 'w', newline='') as f:
            w = csv.DictWriter(f, fieldnames=FIELDS); w.writeheader()
            for k in sorted(self.rows): w.writerow(self.rows[k])
        os.makedirs(LABEL_DIR, exist_ok=True)
        json.dump(self.labels, open(os.path.join(LABEL_DIR, self.tool + '.json'), 'w'), indent=1, ensure_ascii=False)
        def covered(k):  # a label covers its own file, a directory item's files, a session's <id>/ folder, and a db's #rows
            for lk in self.labels:
                base = lk.split('#', 1)[0]
                stem = os.path.splitext(base)[0]
                if k in (lk, base) or k.startswith(base.rstrip('/') + '/') or k.startswith(stem + '/'): return True
            return False
        unlabeled = [k for k in self.rows if not covered(k)]
        print(f'{self.tool}: {len(self.rows)} files, {len(self.labels)} labelled items; files without a label: {len(unlabeled)}')
        for k in unlabeled[:10]: print('   unlabelled:', k)
        return len(self.rows)

# ---------------- native formats ----------------
# blocks for docx/pdf: list of (kind, value): ('title'|'h1'|'h2'|'h3'|'p'|'bullet'|'number'|'quote'|'code', str),
# ('table', [[cells], ...]) (first row = header), ('pagebreak', None)

def _core_props(props, title, author, created, modified):
    props.title = title or ''
    props.author = author; props.last_modified_by = author
    props.created = dt(created).astimezone(UTC).replace(tzinfo=None)
    props.modified = dt(modified or created).astimezone(UTC).replace(tzinfo=None)



def _pin_ooxml(data, created, modified):
    """python-docx/openpyxl/python-pptx stamp the save time into docProps/core.xml and into every zip entry, and add a
    'generated by' description. Pin core dates to created/modified, zip entry times to modified, drop the description,
    so the file is dated where it belongs and identical across runs."""
    import re as _re
    c = dt(created).astimezone(UTC).strftime('%Y-%m-%dT%H:%M:%SZ'); m = dt(modified or created).astimezone(UTC).strftime('%Y-%m-%dT%H:%M:%SZ')
    zin = zipfile.ZipFile(io.BytesIO(data)); out = io.BytesIO()
    stamp = dt(modified or created).timetuple()[:6]
    with zipfile.ZipFile(out, 'w', zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            b = zin.read(info.filename)
            if info.filename == 'docProps/core.xml':
                t = b.decode('utf-8')
                t = _re.sub(r'(<dcterms:created[^>]*>)[^<]*(</dcterms:created>)', lambda x: x.group(1) + c + x.group(2), t)
                t = _re.sub(r'(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)', lambda x: x.group(1) + m + x.group(2), t)
                t = _re.sub(r'<dc:description>[^<]*</dc:description>', '', t)
                t = _re.sub(r'<dc:description/>', '', t)
                b = t.encode('utf-8')
            zi = zipfile.ZipInfo(info.filename, date_time=stamp); zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = info.external_attr
            zout.writestr(zi, b)
    return out.getvalue()


def docx_bytes(blocks, title, author, created, modified):
    import docx
    d = docx.Document()
    for kind, v in blocks:
        if kind == 'title': d.add_heading(v, 0)
        elif kind in ('h1', 'h2', 'h3'): d.add_heading(v, int(kind[1]))
        elif kind == 'bullet': d.add_paragraph(v, style='List Bullet')
        elif kind == 'number': d.add_paragraph(v, style='List Number')
        elif kind == 'quote': d.add_paragraph(v, style='Intense Quote')
        elif kind == 'code':
            para = d.add_paragraph(); r = para.add_run(v); r.font.name = 'Courier New'
        elif kind == 'table':
            t = d.add_table(rows=len(v), cols=max(len(r) for r in v)); t.style = 'Table Grid'
            for i, row in enumerate(v):
                for j, c in enumerate(row): t.cell(i, j).text = str(c)
        elif kind == 'pagebreak': d.add_page_break()
        else: d.add_paragraph(v)
    _core_props(d.core_properties, title, author, created, modified)
    b = io.BytesIO(); d.save(b); return _pin_ooxml(b.getvalue(), created, modified)


def xlsx_bytes(sheets, title, author, created, modified):
    """sheets: {sheet name: [[cell, ...], ...]}; cells may be numbers, strings, dates, or '=FORMULA'."""
    import openpyxl
    wb = openpyxl.Workbook(); wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name[:31])
        for r in rows: ws.append(list(r))
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = min(48, max(9, max(len(str(c.value or '')) for c in col) + 2))
    _core_props(wb.properties, title, author, created, modified)
    wb.properties.creator = author
    b = io.BytesIO(); wb.save(b); return _pin_ooxml(b.getvalue(), created, modified)


def pptx_bytes(slides, title, author, created, modified):
    """slides: [(title, [bullet, ...] | str, notes|None)]"""
    from pptx import Presentation
    pr = Presentation()
    for s in slides:
        t, body = s[0], s[1]; notes = s[2] if len(s) > 2 else None
        sl = pr.slides.add_slide(pr.slide_layouts[1] if body else pr.slide_layouts[0])
        sl.shapes.title.text = t
        if body:
            tf = sl.placeholders[1].text_frame
            items = body if isinstance(body, list) else [body]
            tf.text = items[0]
            for x in items[1:]: tf.add_paragraph().text = x
        if notes: sl.notes_slide.notes_text_frame.text = notes
    _core_props(pr.core_properties, title, author, created, modified)
    b = io.BytesIO(); pr.save(b); return _pin_ooxml(b.getvalue(), created, modified)


def pdf_bytes(blocks, title, author, created, modified, two_column=False):
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, PageBreak, Preformatted
    from xml.sax.saxutils import escape
    ss = getSampleStyleSheet()
    b = io.BytesIO()
    c0, m0 = dt(created), dt(modified or created)

    def stamp(canvas, doc):
        canvas.setCreator('Microsoft Word' if author else 'LaTeX with hyperref'); canvas.setAuthor(author or ''); canvas.setTitle(title or '')
        canvas._doc.info.creationDate = c0.strftime("D:%Y%m%d%H%M%S+00'00'")
        canvas._doc.info.modDate = m0.strftime("D:%Y%m%d%H%M%S+00'00'")
    doc = SimpleDocTemplate(b, pagesize=letter, title=title or '', author=author or '', invariant=1)
    flow = []
    for kind, v in blocks:
        if kind == 'title': flow.append(Paragraph(escape(v), ss['Title']))
        elif kind in ('h1', 'h2', 'h3'): flow.append(Paragraph(escape(v), ss['Heading' + kind[1]]))
        elif kind in ('bullet', 'number'): flow.append(Paragraph('• ' + escape(v), ss['BodyText']))
        elif kind == 'code': flow.append(Preformatted(v, ss['Code']))
        elif kind == 'table': flow.append(Table(v)); flow.append(Spacer(1, 8))
        elif kind == 'pagebreak': flow.append(PageBreak())
        else: flow.append(Paragraph(escape(v), ss['BodyText'])); flow.append(Spacer(1, 4))
    doc.build(flow, onFirstPage=stamp, onLaterPages=stamp)
    # invariant=1 (deterministic output) writes 2000-01-01 for both dates; put the real ones in (same length, so the xref
    # offsets stay valid)
    out = b.getvalue()
    out = out.replace(b"/CreationDate (D:20000101000000+00'00')", b"/CreationDate (" + c0.astimezone(UTC).strftime("D:%Y%m%d%H%M%S+00'00'").encode() + b")", 1)
    out = out.replace(b"/ModDate (D:20000101000000+00'00')", b"/ModDate (" + m0.astimezone(UTC).strftime("D:%Y%m%d%H%M%S+00'00'").encode() + b")", 1)
    return out


def addr(x):
    return email.utils.formataddr(x) if isinstance(x, tuple) else x


def eml_bytes(frm, to, subject, date, body, cc=None, html=None, attachments=(), message_id=None, in_reply_to=None, references=None,
              headers=None, domain=None):
    m = EmailMessage()
    m['From'] = addr(frm)
    m['To'] = ', '.join(addr(x) for x in (to if isinstance(to, list) else [to]))
    if cc: m['Cc'] = ', '.join(addr(x) for x in (cc if isinstance(cc, list) else [cc]))
    m['Subject'] = subject
    m['Date'] = rfc2822(date)
    dom = domain or (addr(frm).split('@')[-1].strip('>') if '@' in addr(frm) else 'mail.gmail.com')
    m['Message-ID'] = message_id or f'<{hexid((subject, iso(date), addr(frm)), 24)}@{dom}>'
    if in_reply_to: m['In-Reply-To'] = in_reply_to
    if references: m['References'] = references if isinstance(references, str) else ' '.join(references)
    m['MIME-Version'] = '1.0'
    for k, v in (headers or {}).items(): m[k] = v
    m.set_content(body)
    if html: m.add_alternative(html, subtype='html')
    for name, data, mime in attachments:
        mt, st = mime.split('/', 1)
        m.add_attachment(data, maintype=mt, subtype=st, filename=name)
    return bytes(m)


def message_id_of(eml_path):
    from email.parser import BytesHeaderParser
    with open(eml_path, 'rb') as f: return BytesHeaderParser().parse(f)['Message-ID']


def git_repo(path, commits, author=('David Barron', 'dbarron410@vt.edu'), branch='main', remote=None):
    """Create a git repo at path. commits: [(iso date, message, {relpath: content or None to delete}), ...] in order;
    commit author/committer dates are backdated. Returns the HEAD sha."""
    os.makedirs(path, exist_ok=True)
    env = dict(os.environ, GIT_CONFIG_GLOBAL='/dev/null', GIT_CONFIG_SYSTEM='/dev/null')
    run = lambda *a, **k: subprocess.run(['git', '-C', path, *a], check=True, capture_output=True, text=True, env={**env, **k.get('extra', {})})
    run('init', '-q', '-b', branch)
    for when, msg, files, *who in commits:
        for rel, content in files.items():
            fp = os.path.join(path, rel)
            if content is None:
                if os.path.exists(fp): os.remove(fp)
                continue
            os.makedirs(os.path.dirname(fp), exist_ok=True)
            with open(fp, 'wb' if isinstance(content, bytes) else 'w') as f: f.write(content)
        run('add', '-A')
        a = who[0] if who else author
        d = dt(when).isoformat()
        run('-c', f'user.name={a[0]}', '-c', f'user.email={a[1]}', 'commit', '-q', '--allow-empty', '-m', msg,
            extra={'GIT_AUTHOR_DATE': d, 'GIT_COMMITTER_DATE': d})
    if remote: run('remote', 'add', 'origin', remote)
    # working-tree file times = last commit that touched them (approximately: last commit date)
    last = dt(commits[-1][0]).timestamp()
    for dp, dn, fn in os.walk(path):
        if '/.git' in dp: continue
        for f in fn: os.utime(os.path.join(dp, f), (last, last))
    return run('rev-parse', 'HEAD').stdout.strip()


SHARED = os.path.join(ROOT, '_scripts', 'synth', 'shared')  # assets several generators copy (the same-paper hard case)
